from pathlib import Path
import fitz
import pytest
from docx import Document

from app.models import Finding
from app.redact import RedactionError, redact_file, redact_pdf, redact_txt


def finding(text, approved=True):
    return Finding(id="test-id", text=text, category="EMAIL", source_layer="regex",
                   reason="synthetic test", location=None, approved=approved)


def test_txt_redacts_and_verifies_saved_copy(tmp_path):
    src, dst = tmp_path / "input.txt", tmp_path / "output.txt"
    src.write_text("Contact morgan@example.invalid.\nAgain morgan@example.invalid.", encoding="utf-8")
    result = redact_txt(src, dst, [finding("morgan@example.invalid")])
    assert result["verified_clean"] is True and result["redacted_occurrences"] == 2
    assert "morgan@example.invalid" not in dst.read_text(encoding="utf-8")
    assert "morgan@example.invalid" in src.read_text(encoding="utf-8")


def test_txt_ignores_unapproved_findings(tmp_path):
    src, dst = tmp_path / "in.txt", tmp_path / "out.txt"
    src.write_text("morgan@example.invalid", encoding="utf-8")
    result = redact_txt(src, dst, [finding("morgan@example.invalid", approved=False)])
    assert result["verified_clean"] is True
    assert dst.read_text(encoding="utf-8") == "morgan@example.invalid"


def test_txt_missing_approved_value_fails_without_output(tmp_path):
    src, dst = tmp_path / "in.txt", tmp_path / "out.txt"
    src.write_text("Nothing sensitive here.", encoding="utf-8")
    with pytest.raises(RedactionError):
        redact_txt(src, dst, [finding("missing@example.invalid")])
    assert not dst.exists()


def test_refuses_to_overwrite_source(tmp_path):
    src = tmp_path / "same.txt"
    src.write_text("morgan@example.invalid", encoding="utf-8")
    with pytest.raises(RedactionError):
        redact_txt(src, src, [finding("morgan@example.invalid")])
    assert "morgan@example.invalid" in src.read_text(encoding="utf-8")


def test_pdf_real_redaction_and_metadata_removal(tmp_path):
    src, dst = tmp_path / "input.pdf", tmp_path / "output.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Contact morgan@example.invalid")
    doc.set_metadata({"author": "FAKE AUTHOR: Morgan Example", "title": "Synthetic"})
    doc.save(src)
    doc.close()
    result = redact_pdf(src, dst, [finding("morgan@example.invalid")])
    assert result["verified_clean"] is True
    with fitz.open(dst) as out:
        assert "morgan@example.invalid" not in "\n".join(p.get_text() for p in out)
        assert "FAKE AUTHOR: Morgan Example" not in str(out.metadata)
    with fitz.open(src) as original:
        assert "morgan@example.invalid" in original[0].get_text()


def test_pdf_missing_value_fails_without_output(tmp_path):
    src, dst = tmp_path / "input.pdf", tmp_path / "output.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Synthetic harmless text")
    doc.save(src)
    doc.close()
    with pytest.raises(RedactionError):
        redact_pdf(src, dst, [finding("missing@example.invalid")])
    assert not dst.exists()


def test_docx_dispatch_redacts_and_verifies(tmp_path):
    src, dst = tmp_path / "input.docx", tmp_path / "output.docx"
    doc = Document()
    doc.add_paragraph("Contact morgan@example.invalid")
    doc.save(src)
    result = redact_file(src, dst, [finding("morgan@example.invalid")])
    assert result["verified_clean"] is True
    assert "morgan@example.invalid" not in Document(dst).paragraphs[0].text
