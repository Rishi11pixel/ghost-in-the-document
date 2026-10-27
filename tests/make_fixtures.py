"""Generate synthetic fixtures for Ghost in the Document. Never use real personal data."""
import json
from pathlib import Path

import fitz
from docx import Document

ROOT = Path(__file__).resolve().parent
FIXTURES = ROOT / "fixtures"
GROUND_TRUTH = ROOT / "ground_truth" / "synthetic_case.json"

def main():
    FIXTURES.mkdir(parents=True, exist_ok=True)
    (ROOT / "ground_truth").mkdir(parents=True, exist_ok=True)

    lines = [
        "SYNTHETIC TEST DOCUMENT — NOT A REAL PERSON",
        "Name: Morgan Example",
        "Email: morgan.example@example.invalid",
        "Phone: +91 90000 12345",
        "Aadhaar-like test ID: 1234 5678 9012",
        "PAN-like test ID: ABCDE1234F",
        "Address: 42 Example Lane, Sample Nagar, Delhi 110001",
    ]
    (FIXTURES / "synthetic_case.txt").write_text("\n".join(lines) + "\n", encoding="utf-8")

    pdf = fitz.open()
    page = pdf.new_page()
    y = 72
    for line in lines:
        page.insert_text((72, y), line, fontsize=11)
        y += 22
    pdf.set_metadata({
        "title": "Synthetic Privacy Test",
        "author": "FAKE AUTHOR: Morgan Example",
        "subject": "Synthetic fixture only",
        "keywords": "FAKE-METADATA-KEYWORD",
        "creator": "Ghost in the Document test fixture",
    })
    pdf.save(FIXTURES / "synthetic_case.pdf")
    pdf.close()

    doc = Document()
    doc.core_properties.title = "Synthetic Privacy Test"
    doc.core_properties.author = "FAKE AUTHOR: Morgan Example"
    doc.core_properties.last_modified_by = "FAKE EDITOR: Taylor Sample"
    doc.core_properties.subject = "Synthetic fixture only"
    doc.sections[0].header.paragraphs[0].text = "Confidential test header: Morgan Example"
    doc.add_heading(lines[0], level=1)
    for line in lines[1:]:
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Field"
    table.rows[0].cells[1].text = "Synthetic value"
    row = table.add_row().cells
    row[0].text = "Contact"
    row[1].text = "morgan.example@example.invalid"
    doc.sections[0].footer.paragraphs[0].text = "Synthetic footer marker: FAKE-FOOTER-001"
    doc.save(FIXTURES / "synthetic_case.docx")

    truth = {
        "fixture_set": "ghost-in-the-document-rung-1",
        "synthetic_data_only": True,
        "notes": [
            "All identities and values are fictional test markers.",
            "The .invalid email domain is reserved for invalid/example use.",
            "Ground truth lists expected spans, not a claim that current code detects them.",
            "DOCX comments and tracked changes will be added in a later fixture update."
        ],
        "expected_findings": [
            {"text": "Morgan Example", "category": "NAME", "locations": ["TXT line 2", "PDF page 1", "DOCX paragraph 2", "DOCX header"]},
            {"text": "morgan.example@example.invalid", "category": "EMAIL", "locations": ["TXT line 3", "PDF page 1", "DOCX paragraph 3", "DOCX table row 2"]},
            {"text": "+91 90000 12345", "category": "PHONE", "locations": ["TXT line 4", "PDF page 1", "DOCX paragraph 4"]},
            {"text": "1234 5678 9012", "category": "ID_NUMBER", "locations": ["TXT line 5", "PDF page 1", "DOCX paragraph 5"]},
            {"text": "ABCDE1234F", "category": "ID_NUMBER", "locations": ["TXT line 6", "PDF page 1", "DOCX paragraph 6"]},
            {"text": "42 Example Lane, Sample Nagar, Delhi 110001", "category": "ADDRESS", "locations": ["TXT line 7", "PDF page 1", "DOCX paragraph 7"]},
            {"text": "FAKE AUTHOR: Morgan Example", "category": "METADATA", "locations": ["PDF metadata author", "DOCX core_properties.author"]},
            {"text": "FAKE EDITOR: Taylor Sample", "category": "METADATA", "locations": ["DOCX core_properties.last_modified_by"]},
            {"text": "FAKE-FOOTER-001", "category": "HIDDEN_TEXT", "locations": ["DOCX footer"]}
        ]
    }
    GROUND_TRUTH.write_text(json.dumps(truth, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    print("Synthetic TXT, PDF, DOCX and ground-truth JSON fixtures created.")

if __name__ == "__main__":
    main()
