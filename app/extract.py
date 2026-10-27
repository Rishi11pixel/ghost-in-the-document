"""Local text extraction for TXT, text-based PDF, and DOCX files.

Scanned/near-empty PDF pages are rejected rather than silently passed through.
This module extracts text only; it does not redact or certify documents as safe.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Literal

import fitz
from docx import Document


@dataclass(frozen=True)
class TextSegment:
    text: str
    location: str


@dataclass(frozen=True)
class ExtractionResult:
    filename: str
    file_type: Literal["txt", "pdf", "docx"]
    text: str
    segments: list[TextSegment]


class UnsupportedFileTypeError(ValueError):
    """Raised when the input file type is not supported."""


class UnsupportedScannedPDFError(ValueError):
    """Raised when one or more PDF pages have too little extractable text."""

    def __init__(self, pages: list[int]):
        self.pages = pages
        page_label = ", ".join(str(page) for page in pages)
        super().__init__(
            f"PDF contains page(s) with insufficient extractable text: {page_label}. "
            "Scanned PDFs/OCR are unsupported; no output was produced."
        )


def _pdf_page_has_extractable_text(text: str, min_alphanumeric_chars: int) -> bool:
    return sum(char.isalnum() for char in text) >= min_alphanumeric_chars


def extract_document(
    path: str | Path,
    *,
    min_pdf_page_chars: int = 20,
) -> ExtractionResult:
    """Extract visible text and basic locations from a supported document.

    PDF policy: if ANY page has fewer than `min_pdf_page_chars` alphanumeric
    characters, reject the whole document. This conservative threshold may
    reject intentionally sparse PDFs; callers can set it explicitly for tests
    or known document types, but should not disable the guard for untrusted files.
    """
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(f"Input file does not exist: {source}")

    suffix = source.suffix.lower()

    if suffix == ".txt":
        content = source.read_text(encoding="utf-8-sig")
        segments = [
            TextSegment(text=line, location=f"line {index}")
            for index, line in enumerate(content.splitlines(), start=1)
            if line.strip()
        ]
        return ExtractionResult(
            filename=source.name,
            file_type="txt",
            text="\n".join(segment.text for segment in segments),
            segments=segments,
        )

    if suffix == ".pdf":
        segments: list[TextSegment] = []
        scanned_pages: list[int] = []
        with fitz.open(source) as pdf:
            if pdf.needs_pass:
                raise ValueError("Password-protected PDFs are unsupported.")
            for page_index, page in enumerate(pdf, start=1):
                page_text = page.get_text("text").strip()
                if not _pdf_page_has_extractable_text(page_text, min_pdf_page_chars):
                    scanned_pages.append(page_index)
                else:
                    segments.append(
                        TextSegment(text=page_text, location=f"page {page_index}")
                    )
        if scanned_pages:
            raise UnsupportedScannedPDFError(scanned_pages)
        return ExtractionResult(
            filename=source.name,
            file_type="pdf",
            text="\n\n".join(segment.text for segment in segments),
            segments=segments,
        )

    if suffix == ".docx":
        document = Document(source)
        segments = []

        for index, paragraph in enumerate(document.paragraphs, start=1):
            if paragraph.text.strip():
                segments.append(
                    TextSegment(text=paragraph.text, location=f"paragraph {index}")
                )

        # Tables are represented separately; repeated/merged cells may appear
        # more than once in complex DOCX files, so deduplication is a later concern.
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                for cell_index, cell in enumerate(row.cells, start=1):
                    if cell.text.strip():
                        segments.append(
                            TextSegment(
                                text=cell.text,
                                location=(
                                    f"table {table_index}, row {row_index}, "
                                    f"cell {cell_index}"
                                ),
                            )
                        )

        # Headers and footers are included in extraction. Hidden XML, comments,
        # tracked changes, footnotes and text boxes require the hidden scanner.
        for section_index, section in enumerate(document.sections, start=1):
            for kind, part in (("header", section.header), ("footer", section.footer)):
                for paragraph_index, paragraph in enumerate(part.paragraphs, start=1):
                    if paragraph.text.strip():
                        segments.append(
                            TextSegment(
                                text=paragraph.text,
                                location=(
                                    f"section {section_index} {kind} "
                                    f"paragraph {paragraph_index}"
                                ),
                            )
                        )
                for table_index, table in enumerate(part.tables, start=1):
                    for row_index, row in enumerate(table.rows, start=1):
                        for cell_index, cell in enumerate(row.cells, start=1):
                            if cell.text.strip():
                                segments.append(
                                    TextSegment(
                                        text=cell.text,
                                        location=(
                                            f"section {section_index} {kind} "
                                            f"table {table_index}, row {row_index}, "
                                            f"cell {cell_index}"
                                        ),
                                    )
                                )

        return ExtractionResult(
            filename=source.name,
            file_type="docx",
            text="\n".join(segment.text for segment in segments),
            segments=segments,
        )

    raise UnsupportedFileTypeError(
        f"Unsupported file extension {suffix!r}. Supported types: .pdf, .docx, .txt"
    )


def supported_extension(filename: str) -> bool:
    """Return whether a filename uses a supported extension."""
    return filename.lower().endswith((".pdf", ".docx", ".txt"))
